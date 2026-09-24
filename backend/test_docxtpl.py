import docx
from docxtpl import DocxTemplate

doc = docx.Document()
t = doc.add_table(rows=3, cols=2)

# Row 0: {%tr for item in items %}
t.rows[0].cells[0].paragraphs[0].text = "{%tr for item in items %}"

# Row 1: The actual data row to repeat
t.rows[1].cells[0].paragraphs[0].text = "{{ item.a }}"
t.rows[1].cells[1].paragraphs[0].text = "{{ item.b }}"

# Row 2: {%tr endfor %}
t.rows[2].cells[0].paragraphs[0].text = "{%tr endfor %}"

doc.save("test_template.docx")

tpl = DocxTemplate("test_template.docx")
tpl.render({'items': [{'a': 1, 'b': 2}, {'a': 3, 'b': 4}]})
tpl.save("test_result.docx")
print("SUCCESS")
