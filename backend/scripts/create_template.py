import os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

def create():
    doc = Document()
    
    # 1. Top right text
    p_top_right = doc.add_paragraph()
    p_top_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p_top_right.add_run("Mẫu PC08\nBH theo Nghị định số 105/2025/NĐ-CP\nNgày 15/5/2025")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.bold = True
    
    doc.add_paragraph() # Spacer
    
    # 2. Top center text (Cơ quan chủ quản)
    p_cq = doc.add_paragraph()
    p_cq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_cq1 = p_cq.add_run("CÔNG AN TỈNH PHÚ THỌ\n")
    run_cq1.font.name = 'Times New Roman'
    run_cq1.font.size = Pt(13)
    
    run_cq2 = p_cq.add_run("PHÒNG CẢNH SÁT PHÒNG CHÁY, CHỮA CHÁY VÀ CỨU NẠN, CỨU HỘ\n")
    run_cq2.font.name = 'Times New Roman'
    run_cq2.font.size = Pt(13)
    run_cq2.bold = True
    
    run_cq3 = p_cq.add_run("_______________\n\n")
    run_cq3.font.name = 'Times New Roman'
    
    run_cq4 = p_cq.add_run("Số: {{ so_phuong_an }}")
    run_cq4.font.name = 'Times New Roman'
    run_cq4.font.size = Pt(13)
    
    for _ in range(4): doc.add_paragraph() # Spacers
    
    # 3. Main Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("PHƯƠNG ÁN CHỮA CHÁY, CỨU NẠN, CỨU HỘ")
    run_title.font.name = 'Times New Roman'
    run_title.font.size = Pt(16)
    run_title.bold = True
    
    for _ in range(3): doc.add_paragraph() # Spacers
    
    # 4. Facility Info block
    p_info = doc.add_paragraph()
    p_info.paragraph_format.left_indent = Pt(100) # Indent roughly to match image
    
    def add_info_line(p, label, value, bold_label=True):
        r1 = p.add_run(label)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(13)
        r1.bold = bold_label
        r2 = p.add_run(value + "\n")
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(13)
        r2.bold = False
        
    add_info_line(p_info, "Tên cơ sở: ", "{{ ten_co_so }}")
    add_info_line(p_info, "Địa chỉ: ", "{{ dia_chi_co_so }}")
    add_info_line(p_info, "Điện thoại: ", "{{ dien_thoai_co_so }}")
    add_info_line(p_info, "Đơn vị được phân công thực hiện nhiệm vụ chữa cháy, cứu nạn, cứu hộ:\n", "")
    
    r_donvi = p_info.add_run("{{ don_vi_pccc }}\n")
    r_donvi.font.name = 'Times New Roman'
    r_donvi.font.size = Pt(13)
    
    add_info_line(p_info, "Điện thoại: ", "114")
    
    for _ in range(8): doc.add_paragraph() # Spacers to bottom
    
    # 5. Year at bottom
    p_year = doc.add_paragraph()
    p_year.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_year = p_year.add_run("Năm 2026")
    run_year.font.name = 'Times New Roman'
    run_year.font.size = Pt(14)
    run_year.bold = True
    
    doc.add_page_break()
    
    # --- END COVER PAGE ---
    
    doc.add_heading("A. ĐẶC ĐIỂM CÓ LIÊN QUAN", level=2)
    doc.add_heading("I. VỊ TRÍ", level=3)
    doc.add_paragraph("Tọa độ: {{ toa_do }}")
    doc.add_paragraph("Đông: {{ giap_dong }}\nTây: {{ giap_tay }}\nNam: {{ giap_nam }}\nBắc: {{ giap_bac }}")
    
    doc.add_heading("II. GIAO THÔNG", level=3)
    doc.add_paragraph("Cổng: {{ cong_chinh_rong }}m x {{ cong_chinh_cao }}m")
    doc.add_paragraph("Đường nội bộ: {{ duong_noi_bo_rong }}m, tải trọng {{ duong_noi_bo_tai_trong }} tấn")
    doc.add_paragraph("Giao thông & Nguồn nước: {{ giao_thong_nguon_nuoc }}")
    
    doc.add_heading("III. NGUỒN NƯỚC", level=3)
    t = doc.add_table(rows=4, cols=5)
    t.style = 'Table Grid'
    t.rows[0].cells[0].text = "STT"
    t.rows[0].cells[1].text = "Loại nguồn"
    t.rows[0].cells[2].text = "Trữ lượng"
    t.rows[0].cells[3].text = "Vị trí"
    t.rows[0].cells[4].text = "Ghi chú"
    t.rows[1].cells[0].text = "{%tr for item in danh_sach_nguon_nuoc %}"
    t.rows[2].cells[0].text = "{{ item.stt }}"
    t.rows[2].cells[1].text = "{{ item.loai_nguon }}"
    t.rows[2].cells[2].text = "{{ item.tru_luong }}"
    t.rows[2].cells[3].text = "{{ item.vi_tri }}"
    t.rows[2].cells[4].text = "{{ item.ghi_chu }}"
    t.rows[3].cells[0].text = "{%tr endfor %}"
    
    doc.add_heading("IV. ĐẶC ĐIỂM KIẾN TRÚC & NGUY HIỂM", level=3)
    doc.add_paragraph("Tổng diện tích: {{ tong_dien_tich_xay_dung }} m2")
    doc.add_paragraph("{% for hm in hang_muc_cong_trinh %}- {{ hm.ten }}, {{ hm.dien_tich }}m2, {{ hm.cong_nang }}\n{% endfor %}")
    doc.add_paragraph("Đánh giá nguy hiểm: {{ muc_iv_nguy_hiem_chay_no }}")
    
    doc.add_heading("V. & VI. LỰC LƯỢNG VÀ PHƯƠNG TIỆN", level=3)
    doc.add_paragraph("Người đứng đầu: {{ nguoi_dung_dau }} - {{ sdt_nguoi_dung_dau }}")
    doc.add_paragraph("Đội trưởng PCCC: {{ doi_truong_pccc }} - {{ sdt_doi_truong }} (Tổng: {{ tong_doi_vien_pccc }} đội viên)")
    doc.add_paragraph("Trực trong giờ: {{ truc_trong_gio }} | Ngoài giờ: {{ truc_ngoai_gio }}")
    
    t2 = doc.add_table(rows=4, cols=2)
    t2.style = 'Table Grid'
    t2.rows[0].cells[0].text = "Tên phương tiện"
    t2.rows[0].cells[1].text = "Số lượng"
    t2.rows[1].cells[0].text = "{%tr for pt in phuong_tien_tai_cho %}"
    t2.rows[2].cells[0].text = "{{ pt.ten }}"
    t2.rows[2].cells[1].text = "{{ pt.so_luong }}"
    t2.rows[3].cells[0].text = "{%tr endfor %}"
    
    t3 = doc.add_table(rows=4, cols=2)
    t3.style = 'Table Grid'
    t3.rows[0].cells[0].text = "Hệ thống PCCC"
    t3.rows[0].cells[1].text = "Số lượng"
    t3.rows[1].cells[0].text = "{%tr for ht in he_thong_pccc %}"
    t3.rows[2].cells[0].text = "{{ ht.ten }}"
    t3.rows[2].cells[1].text = "{{ ht.so_luong }}"
    t3.rows[3].cells[0].text = "{%tr endfor %}"

    doc.add_heading("B. PHƯƠNG ÁN XỬ LÝ TÌNH HUỐNG PHỨC TẠP NHẤT", level=2)
    doc.add_paragraph("Thời gian: {{ b_thoi_gian_chay }}")
    doc.add_paragraph("Điểm cháy: {{ b_diem_chay }}")
    doc.add_paragraph("Nguyên nhân: {{ b_nguyen_nhan }}")
    doc.add_paragraph("Chất cháy: {{ b_chat_chay }}")
    doc.add_paragraph("Diện tích cháy: {{ b_dien_tich_chay }} m2")
    doc.add_paragraph("Số người bị nạn: {{ b_so_nguoi_bi_nan }}")
    
    doc.add_heading("Tính toán kỹ thuật", level=3)
    doc.add_paragraph("Giải thích phép toán:\n{{ giai_thich_phep_toan }}")
    doc.add_paragraph("Diện tích chữa cháy: {{ b_dien_tich_chua_chay }} m2")
    doc.add_paragraph("Lưu lượng chữa cháy: {{ b_luu_luong_chua_chay }} l/s")
    doc.add_paragraph("Số lăng chữa cháy: {{ b_so_lang_chua_chay }} | Làm mát: {{ b_so_lang_lam_mat }}")
    doc.add_paragraph("Số xe chữa cháy: {{ b_so_xe_chua_chay }}")
    doc.add_paragraph("Lượng nước (m3): {{ b_luong_nuoc_m3 }}")
    
    doc.add_heading("Chiến thuật và Lưu ý", level=3)
    doc.add_paragraph("Biện pháp: {{ b_bien_phap_chua_chay }}")
    doc.add_paragraph("Phương pháp: {{ b_phuong_phap_chua_chay }}")
    doc.add_paragraph("An toàn lực lượng: {{ b_an_toan_luc_luong }}")
    doc.add_paragraph("Thông tin chỉ huy: {{ b_thong_tin_chi_huy }}")
    doc.add_paragraph("Chất kỵ nước: {{ b_chat_ky_nuoc }}")
    doc.add_paragraph("Khu vực khói độc: {{ b_khu_vuc_khoi_doc }}")
    doc.add_paragraph("Khu vực khó tiếp cận: {{ b_khu_vuc_kho_tiep_can }}")
    doc.add_paragraph("Nguy cơ nổ: {{ b_nguy_co_no }}")
    
    doc.add_heading("Bảng thống kê lực lượng, phương tiện dự kiến huy động", level=3)
    t4 = doc.add_table(rows=4, cols=6)
    t4.style = 'Table Grid'
    t4.rows[0].cells[0].text = "STT"
    t4.rows[0].cells[1].text = "Đơn vị được huy động"
    t4.rows[0].cells[2].text = "Điện thoại"
    t4.rows[0].cells[3].text = "Số người"
    t4.rows[0].cells[4].text = "Phương tiện"
    t4.rows[0].cells[5].text = "Ghi chú"
    t4.rows[1].cells[0].text = "{%tr for bk in bang_thong_ke_luc_luong %}"
    t4.rows[2].cells[0].text = "{{ bk.stt }}"
    t4.rows[2].cells[1].text = "{{ bk.don_vi }}"
    t4.rows[2].cells[2].text = "{{ bk.dien_thoai }}"
    t4.rows[2].cells[3].text = "{{ bk.so_nguoi }}"
    t4.rows[2].cells[4].text = "{{ bk.phuong_tien }}"
    t4.rows[2].cells[5].text = "{{ bk.ghi_chu }}"
    t4.rows[3].cells[0].text = "{%tr endfor %}"

    doc.add_heading("Nhiệm vụ tác chiến", level=3)
    doc.add_paragraph("Lực lượng tại chỗ: {{ b_nhiem_vu_tai_cho }}")
    doc.add_paragraph("Trinh sát: {{ b_nhiem_vu_trinh_sat }}")
    doc.add_paragraph("Xe C1: {{ b_xe_c1 }}")
    doc.add_paragraph("Xe C2: {{ b_xe_c2 }}")
    doc.add_paragraph("Xe C3/CNCH: {{ b_xe_c3_cnch }}")
    doc.add_paragraph("Công an xã: {{ b_nhiem_vu_ca_xa }}")
    doc.add_paragraph("CSGT: {{ b_nhiem_vu_csgt }}")
    doc.add_paragraph("Y tế: {{ b_nhiem_vu_y_te }}")
    doc.add_paragraph("Điện lực: {{ b_nhiem_vu_dien_luc }}")
    
    doc.add_heading("C. TÌNH HUỐNG KHÁC", level=2)
    doc.add_paragraph("1. Tình huống cháy 1: {{ c_tinh_huong_1 }}")
    doc.add_paragraph("2. Tình huống cháy 2: {{ c_tinh_huong_2 }}")
    doc.add_paragraph("3. CNCH - Tên sự cố: {{ c_cnch_ten }}")
    doc.add_paragraph("   Giả định: {{ c_cnch_gia_dinh }}")
    doc.add_paragraph("   Biện pháp xử lý: {{ c_cnch_bien_phap }}")
    
    os.makedirs(os.path.join(os.path.dirname(__file__), '../app/templates'), exist_ok=True)
    path = os.path.join(os.path.dirname(__file__), '../app/templates/template_pc08.docx')
    doc.save(path)
    print(f"Generated template at {path}")

if __name__ == '__main__':
    create()
